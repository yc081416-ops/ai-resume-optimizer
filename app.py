# -*- coding: utf-8 -*-
"""
AI 简历优化助手
----------------
用户上传简历（PDF / DOCX / TXT）+ 粘贴目标岗位 JD，调用 DeepSeek API
输出结构化的简历优化建议。

运行方式：streamlit run app.py

注意：本文件严禁硬编码任何 API Key。
API Key 的读取优先级：Streamlit secrets > 环境变量 > 侧边栏手动输入。
"""

import io
import os
from datetime import datetime

import streamlit as st

# ============================================================
# 常量配置
# ============================================================
APP_TITLE = "AI 简历优化助手"
MAX_FILE_SIZE_MB = 5
MAX_FILE_SIZE = MAX_FILE_SIZE_MB * 1024 * 1024
SUPPORTED_TYPES = ["pdf", "docx", "txt"]

DEEPSEEK_BASE_URL = "https://api.deepseek.com"
DEEPSEEK_MODEL = "deepseek-chat"
DEEPSEEK_TIMEOUT = 180.0

# 需求文档第 5 节的 Prompt 模板（结构原样使用）
PROMPT_TEMPLATE = """你是一位资深招聘顾问，擅长简历优化。请对比下面提供的【简历】和【目标岗位JD】，输出结构化建议。

【简历】
{resume_text}

【目标岗位JD】
{jd_text}

请按以下结构输出（中文，Markdown格式）：

## 1. 匹配度评分
给出 0-100 的总体评分，并用 2-3 句话说明理由。

## 2. 缺失关键词
列出 JD 中出现但简历中缺失/薄弱的 5-8 个关键词，按重要性排序。

## 3. 经历优化建议
针对简历中的每段经历，指出问题（空话/缺数字/缺结果/与岗位无关），并给出改写方向。每条建议格式：**原文问题** → **改写建议**。

## 4. 改写示例
挑选简历中最需要改的一段经历，给出完整改写版本（保留真实信息，只优化表达；不要编造数据或经历）。

## 5. 投递建议
基于匹配度，给出是否建议投递 + 投递前必须改的 Top 3。
"""

SYSTEM_PROMPT = (
    "你是一位资深招聘顾问，服务对象是正在找实习的大学生。"
    "你的建议必须具体、可操作，严格基于用户提供的简历原文，"
    "严禁编造简历中不存在的数字、项目或经历。"
)


# ============================================================
# 简历文件解析
# ============================================================
def _parse_pdf(file_bytes: bytes) -> str:
    """解析 PDF 文件为纯文本。"""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _parse_docx(file_bytes: bytes) -> str:
    """解析 DOCX 文件为纯文本（含表格内容）。"""
    import docx

    document = docx.Document(io.BytesIO(file_bytes))
    blocks = [p.text for p in document.paragraphs]
    # 简历里的信息经常放在表格中，一并提取
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            blocks.append("\t".join(cells))
    return "\n".join(blocks)


def _parse_txt(file_bytes: bytes) -> str:
    """解析 TXT 文件为纯文本，自动尝试常见中文编码。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "utf-16"):
        try:
            return file_bytes.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ValueError("无法识别文本编码")


def parse_resume(uploaded_file) -> str:
    """
    根据扩展名解析上传的简历文件，返回纯文本。

    解析失败时抛出 ValueError，由上层统一转为中文提示。
    """
    name = uploaded_file.name or ""
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""

    if ext not in SUPPORTED_TYPES:
        raise ValueError("不支持的文件格式，请上传 PDF / DOCX / TXT 文件")

    file_bytes = uploaded_file.getvalue()
    if not file_bytes:
        raise ValueError("文件内容为空")

    if ext == "pdf":
        text = _parse_pdf(file_bytes)
    elif ext == "docx":
        text = _parse_docx(file_bytes)
    else:
        text = _parse_txt(file_bytes)

    text = text.replace("\x00", "").strip()
    if not text:
        raise ValueError("未能从文件中提取到文字，可能是扫描件或空文件")
    return text


def extract_resume_text(uploaded_file) -> tuple[bool, str]:
    """
    安全解析入口。

    返回 (是否成功, 文本或错误信息)。
    """
    try:
        return True, parse_resume(uploaded_file)
    except Exception:
        return False, "文件解析失败，请检查文件格式"


# ============================================================
# API Key 读取（优先级：secrets > 环境变量 > 手动输入）
# ============================================================
def _valid(key) -> bool:
    return bool(key) and isinstance(key, str) and key.strip() != "" and not key.strip().startswith("sk-xxx")


def _from_secrets() -> str:
    """
    从 Streamlit secrets 读取 Key。

    注意：直接访问 st.secrets 时，若本地不存在 secrets.toml，
    Streamlit 会在页面上渲染一段英文错误提示。因此这里先用官方的
    load_if_toml_exists() 探测（它不会报错），再取值。
    """
    secrets = getattr(st, "secrets", None)
    if secrets is None:
        return ""
    try:
        if not secrets.load_if_toml_exists():
            # 没有任何 secrets.toml，静默跳过，不打扰用户
            return ""
        return secrets.get("DEEPSEEK_API_KEY", "") or ""
    except Exception:
        return ""


def resolve_api_key(manual_key: str = "") -> tuple[str, str]:
    """
    按优先级获取 API Key。

    返回 (key, 来源说明)。key 为空字符串表示未配置。
    """
    key = _from_secrets()
    if _valid(key):
        return key.strip(), "Streamlit secrets"

    key = os.getenv("DEEPSEEK_API_KEY", "")
    if _valid(key):
        return key.strip(), "环境变量"

    if _valid(manual_key):
        return manual_key.strip(), "侧边栏手动输入"

    return "", ""


# ============================================================
# DeepSeek 调用
# ============================================================
def build_prompt(resume_text: str, jd_text: str) -> str:
    """把简历文本和 JD 文本填入 Prompt 模板。"""
    return PROMPT_TEMPLATE.replace("{resume_text}", resume_text).replace("{jd_text}", jd_text)


def _friendly_error(exc: Exception) -> str:
    """把 SDK 异常翻译成中文提示。"""
    try:
        from openai import (
            APIConnectionError,
            APIStatusError,
            APITimeoutError,
            AuthenticationError,
            BadRequestError,
            NotFoundError,
            PermissionDeniedError,
            RateLimitError,
        )
    except Exception:
        return "AI 分析失败：{}".format(exc)

    # 注意：APITimeoutError 是 APIConnectionError 的子类，需先判断
    if isinstance(exc, APITimeoutError):
        return "请求超时：DeepSeek 接口响应过慢，请稍后重试。"
    if isinstance(exc, AuthenticationError):
        return "API Key 无效或已过期，请检查 DEEPSEEK_API_KEY 是否正确。"
    if isinstance(exc, PermissionDeniedError):
        return "API Key 无访问权限，请确认账号状态或更换 Key。"
    if isinstance(exc, RateLimitError):
        return "请求过于频繁或账户余额不足，请稍后重试或检查 DeepSeek 账户余额。"
    if isinstance(exc, NotFoundError):
        return "接口地址或模型不存在，请确认 base_url 与模型名称（deepseek-chat）正确。"
    if isinstance(exc, BadRequestError):
        return "请求参数有误，可能是简历或 JD 内容过长，请精简后重试。"
    if isinstance(exc, APIConnectionError):
        return "网络连接失败：无法访问 DeepSeek 接口，请检查网络或代理设置。"
    if isinstance(exc, APIStatusError):
        return "接口返回错误（HTTP {}），请稍后重试。".format(getattr(exc, "status_code", "未知"))
    return "AI 分析失败：{}".format(exc)


def call_deepseek(api_key: str, prompt: str) -> tuple[bool, str]:
    """
    调用 DeepSeek Chat API。

    返回 (是否成功, 结果 Markdown 或中文错误提示)。
    """
    try:
        from openai import OpenAI
    except Exception:
        return False, "缺少 openai 依赖，请先执行 pip install -r requirements.txt"

    try:
        client = OpenAI(api_key=api_key, base_url=DEEPSEEK_BASE_URL, timeout=DEEPSEEK_TIMEOUT)
        response = client.chat.completions.create(
            model=DEEPSEEK_MODEL,
            messages=[
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": prompt},
            ],
            temperature=0.3,
            stream=False,
        )
        content = (response.choices[0].message.content or "").strip()
        if not content:
            return False, "AI 未返回内容，请稍后重试。"
        return True, content
    except Exception as exc:  # noqa: BLE001 - 统一转为中文提示，保证页面不崩溃
        return False, _friendly_error(exc)


# ============================================================
# 页面
# ============================================================
def init_state() -> None:
    st.session_state.setdefault("resume_text", "")
    st.session_state.setdefault("resume_msg", None)
    st.session_state.setdefault("weak_resume", False)
    st.session_state.setdefault("result", "")
    st.session_state.setdefault("result_error", "")
    st.session_state.setdefault("analyzed_at", "")


def render_sidebar() -> str:
    """渲染侧边栏，返回手动输入的 API Key。"""
    with st.sidebar:
        st.header("配置")

        auto_key, source = resolve_api_key()
        if auto_key:
            st.success("已检测到 API Key（来源：{}）".format(source))
        else:
            st.warning("未检测到 API Key，请在下方手动填写。")

        manual_key = st.text_input(
            "DeepSeek API Key",
            type="password",
            placeholder="sk-xxxxxxxx",
            help="仅在未配置 secrets / 环境变量时需要填写，Key 只保存在当前会话内存中。",
        )

        st.divider()
        st.caption("API Key 读取优先级：")
        st.caption("1. Streamlit secrets（.streamlit/secrets.toml）")
        st.caption("2. 环境变量 DEEPSEEK_API_KEY")
        st.caption("3. 本输入框手动填写")
        st.divider()
        st.caption("支持格式：PDF / DOCX / TXT，单个文件 ≤ {} MB".format(MAX_FILE_SIZE_MB))
        st.caption("模型：{}".format(DEEPSEEK_MODEL))
    return manual_key


def render_input_column() -> tuple[object, str]:
    """渲染左侧输入区，返回 (上传文件, JD 文本)。"""
    st.subheader("① 输入简历与岗位")

    uploaded_file = st.file_uploader(
        "上传简历文件（PDF / DOCX / TXT，≤ {} MB）".format(MAX_FILE_SIZE_MB),
        type=SUPPORTED_TYPES,
        accept_multiple_files=False,
    )

    if uploaded_file is not None:
        size = getattr(uploaded_file, "size", None)
        if size is None:
            try:
                size = len(uploaded_file.getvalue())
            except Exception:
                size = 0
        if size and size > MAX_FILE_SIZE:
            st.session_state.resume_text = ""
            st.session_state.resume_msg = None
            st.error("文件过大（{:.1f} MB），请上传不超过 {} MB 的文件。".format(size / 1024 / 1024, MAX_FILE_SIZE_MB))
        else:
            signature = (uploaded_file.name, size)
            if st.session_state.get("_resume_signature") != signature:
                ok, text_or_error = extract_resume_text(uploaded_file)
                st.session_state["_resume_signature"] = signature
                if ok:
                    st.session_state.resume_text = text_or_error
                    st.session_state.weak_resume = len(text_or_error) < 200
                    st.session_state.resume_msg = "已成功读取 {} 字".format(len(text_or_error))
                else:
                    st.session_state.resume_text = ""
                    st.session_state.resume_msg = text_or_error

    if st.session_state.resume_text:
        st.success("已成功读取 {} 字".format(len(st.session_state.resume_text)))
        with st.expander("查看解析出的简历文本"):
            if st.session_state.weak_resume:
                st.warning("提示：文本内容较少，可能是扫描版 PDF（图片型），建议改用 DOCX / TXT 或可复制文字的 PDF。")
            st.text_area("简历原文", st.session_state.resume_text, height=220, disabled=True, label_visibility="collapsed")
    elif st.session_state.resume_msg:
        st.error("文件解析失败，请检查文件格式")
        with st.expander("失败详情"):
            st.write(st.session_state.resume_msg)

    st.write("")
    jd_text = st.text_area(
        "目标岗位 JD",
        placeholder="在此粘贴目标岗位的职位描述（JD）",
        height=240,
        key="jd_text",
    )
    st.caption("JD 字数：{} 字".format(len(jd_text.strip()) if jd_text else 0))

    return uploaded_file, jd_text


def render_output_column(manual_key: str) -> None:
    """渲染右侧输出区（分析按钮 + 结果展示）。"""
    st.subheader("② 分析结果")

    resume_text = st.session_state.resume_text
    jd_text = st.session_state.get("jd_text", "").strip()

    clicked = st.button("开始分析", type="primary", use_container_width=True)

    if clicked:
        st.session_state.result = ""
        st.session_state.result_error = ""

        if not resume_text:
            st.session_state.result_error = "请先上传并成功解析一份简历（支持 PDF / DOCX / TXT，≤ {} MB）。".format(MAX_FILE_SIZE_MB)
        elif not jd_text:
            st.session_state.result_error = "请先粘贴目标岗位的 JD。"
        else:
            api_key, _source = resolve_api_key(manual_key)
            if not api_key:
                st.session_state.result_error = (
                    "未配置 API Key。请通过以下任一方式配置：\n\n"
                    "1. 在 `.streamlit/secrets.toml` 中设置 `DEEPSEEK_API_KEY`\n"
                    "2. 设置环境变量 `DEEPSEEK_API_KEY`\n"
                    "3. 在左侧侧边栏手动输入 API Key"
                )
            else:
                prompt = build_prompt(resume_text, jd_text)
                with st.spinner("AI 正在分析，请稍候…"):
                    ok, payload = call_deepseek(api_key, prompt)
                if ok:
                    st.session_state.result = payload
                    st.session_state.analyzed_at = datetime.now().strftime("%Y%m%d_%H%M%S")
                else:
                    st.session_state.result_error = payload

    if st.session_state.result_error:
        st.error("分析未完成：" + st.session_state.result_error.split("\n")[0])
        if "\n" in st.session_state.result_error:
            with st.expander("查看详细说明"):
                st.markdown(st.session_state.result_error)

    if st.session_state.result:
        st.markdown(st.session_state.result)

        st.divider()
        st.download_button(
            label="下载为 Markdown 文件",
            data=st.session_state.result.encode("utf-8"),
            file_name="简历优化建议_{}.md".format(st.session_state.analyzed_at or "result"),
            mime="text/markdown",
            use_container_width=True,
        )
        st.caption("也可以直接在上方结果区域全选复制。")
    else:
        if not st.session_state.result_error:
            st.info("上传简历并粘贴 JD 后，点击「开始分析」即可生成优化建议。")


def main() -> None:
    st.set_page_config(page_title=APP_TITLE, page_icon="📄", layout="wide")
    init_state()

    st.title(APP_TITLE)
    st.caption("上传简历 + 粘贴目标岗位 JD，一键获得结构化优化建议。")

    manual_key = render_sidebar()

    left, right = st.columns([1, 1], gap="large")
    with left:
        render_input_column()
    with right:
        render_output_column(manual_key)


if __name__ == "__main__":
    main()
